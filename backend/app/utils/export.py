"""
Utilitaire d'export CSV / Excel (XLSX) partagé par les logs et les alertes.

Version esthétisée : bannière de titre, en-tête coloré, lignes zébrées,
coloration automatique selon la sévérité/criticité, bordures, colonnes
triées dans un ordre logique (au lieu de l'ordre alphabétique brut).
"""
import csv
import io
from datetime import datetime
from typing import List, Dict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from fastapi.responses import StreamingResponse

# ─────────────────────────────────────────────────────────────────────────
# Ordre logique des colonnes (les colonnes non listées sont ajoutées après,
# triées alphabétiquement). Couvre à la fois les logs et les alertes.
# ─────────────────────────────────────────────────────────────────────────
COLUMN_PRIORITY = [
    "id", "timestamp", "niveau_criticite", "severity", "category", "categorie",
    "statut", "status", "regle_id", "description", "host", "cible_host",
    "source_ip", "destination_ip", "log_type", "perimetre_id", "is_suspect",
    "raw_message",
]

# Libellés FR plus lisibles pour les en-têtes de colonnes connues
COLUMN_LABELS = {
    "id": "ID", "timestamp": "Horodatage", "niveau_criticite": "Criticité",
    "severity": "Sévérité", "category": "Catégorie", "categorie": "Catégorie",
    "statut": "Statut", "status": "Statut", "regle_id": "Règle Déclencheuse",
    "description": "Description", "host": "Hôte", "cible_host": "Hôte Cible",
    "source_ip": "IP Source", "destination_ip": "IP Destination",
    "log_type": "Type de Log", "perimetre_id": "Périmètre",
    "is_suspect": "Marqué Suspect", "raw_message": "Message Brut",
}

# Couleurs de fond (correspondance sévérité/criticité -> couleur claire)
SEVERITY_COLORS = {
    "CRITICAL": "FEE2E2", "HIGH": "FFEDD5", "WARNING": "FEF9C3",
    "MEDIUM": "FEF9C3", "INFO": "DCFCE7", "LOW": "DCFCE7",
}
SEVERITY_TEXT_COLORS = {
    "CRITICAL": "991B1B", "HIGH": "9A3412", "WARNING": "854D0E",
    "MEDIUM": "854D0E", "INFO": "166534", "LOW": "166534",
}

HEADER_BG = "1F2937"
BANNER_BG = "0F172A"
BORDER_COLOR = "E2E8F0"


def _ordered_columns(rows: List[Dict]) -> List[str]:
    present = {k for row in rows for k in row.keys()}
    ordered = [c for c in COLUMN_PRIORITY if c in present]
    remaining = sorted(present - set(ordered))
    return ordered + remaining


def _flatten_rows(rows: List[Dict]) -> List[Dict]:
    """Aplatit les valeurs non scalaires et rend les booléens lisibles en français."""
    flat = []
    for row in rows:
        flat_row = {}
        for k, v in row.items():
            if isinstance(v, bool):
                flat_row[k] = "Oui" if v else "Non"
            elif isinstance(v, (list, dict)):
                flat_row[k] = str(v)
            else:
                flat_row[k] = v
        flat.append(flat_row)
    return flat


def _format_timestamp(value):
    if not value or not isinstance(value, str):
        return value
    try:
        dt = datetime.fromisoformat(value.replace("Z", "+00:00"))
        return dt.strftime("%d/%m/%Y %H:%M:%S")
    except Exception:
        return value


# ─────────────────────────────────────────────────────────────────────────
# CSV — colonnes réordonnées, en-têtes FR, timestamp lisible
# ─────────────────────────────────────────────────────────────────────────

def export_to_csv(rows: List[Dict], filename: str) -> StreamingResponse:
    rows = _flatten_rows(rows)
    columns = _ordered_columns(rows)

    buffer = io.StringIO()
    writer = csv.DictWriter(
        buffer,
        fieldnames=columns,
        extrasaction="ignore",
        delimiter=";",  # 🟢 point-virgule : Excel FR ouvre le CSV proprement en colonnes sans import manuel
    )
    writer.writerow({c: COLUMN_LABELS.get(c, c.replace("_", " ").title()) for c in columns})
    for row in rows:
        if "timestamp" in row:
            row["timestamp"] = _format_timestamp(row["timestamp"])
        writer.writerow(row)

    buffer.seek(0)
    byte_buffer = io.BytesIO(buffer.getvalue().encode("utf-8-sig"))  # BOM pour Excel/accents FR
    return StreamingResponse(
        byte_buffer,
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ─────────────────────────────────────────────────────────────────────────
# EXCEL — bannière de titre, en-tête coloré, zébrage, coloration par sévérité
# ─────────────────────────────────────────────────────────────────────────

def export_to_excel(rows: List[Dict], filename: str, sheet_title: str = "Export") -> StreamingResponse:
    rows = _flatten_rows(rows)
    columns = _ordered_columns(rows)
    n_cols = max(len(columns), 1)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    thin_border = Border(*(Side(style="thin", color=BORDER_COLOR) for _ in range(4)))

    # ── Ligne 1-2 : bannière de titre ──────────────────────────────────
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    title_cell = ws.cell(row=1, column=1, value=f"SMART SIEM — {sheet_title}")
    title_cell.font = Font(color="FFFFFF", bold=True, size=14)
    title_cell.fill = PatternFill(start_color=BANNER_BG, end_color=BANNER_BG, fill_type="solid")
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
    subtitle_cell = ws.cell(
        row=2, column=1,
        value=f"Généré le {datetime.utcnow().strftime('%d/%m/%Y à %H:%M UTC')}  ·  {len(rows)} ligne(s)",
    )
    subtitle_cell.font = Font(color="CBD5E1", size=10, italic=True)
    subtitle_cell.fill = PatternFill(start_color=BANNER_BG, end_color=BANNER_BG, fill_type="solid")
    subtitle_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 20

    # Ligne 3 vide pour respirer
    ws.row_dimensions[3].height = 6

    # ── Ligne 4 : en-tête des colonnes ─────────────────────────────────
    header_row = 4
    header_fill = PatternFill(start_color=HEADER_BG, end_color=HEADER_BG, fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    for col_idx, col in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=COLUMN_LABELS.get(col, col.replace("_", " ").title()))
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[header_row].height = 22

    # ── Données : zébrage + coloration sévérité + bordures ────────────
    severity_col_idx = next((i for i, c in enumerate(columns, start=1) if c in ("severity", "niveau_criticite")), None)
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for r, row in enumerate(rows, start=header_row + 1):
        if "timestamp" in row:
            row["timestamp"] = _format_timestamp(row["timestamp"])

        severity_value = None
        if severity_col_idx:
            raw_sev = row.get(columns[severity_col_idx - 1], "")
            severity_value = str(raw_sev).upper().strip()

        for col_idx, col in enumerate(columns, start=1):
            cell = ws.cell(row=r, column=col_idx, value=row.get(col, ""))
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=False)

            if severity_value and severity_value in SEVERITY_COLORS:
                cell.fill = PatternFill(start_color=SEVERITY_COLORS[severity_value], end_color=SEVERITY_COLORS[severity_value], fill_type="solid")
                cell.font = Font(color=SEVERITY_TEXT_COLORS[severity_value], bold=(col_idx == severity_col_idx))
            elif (r - header_row) % 2 == 0:
                cell.fill = zebra_fill

    # ── Largeur des colonnes ────────────────────────────────────────────
    for i, col in enumerate(columns, start=1):
        header_len = len(COLUMN_LABELS.get(col, col))
        max_len = max([header_len] + [len(str(row.get(col, ""))) for row in rows]) if rows else header_len
        ws.column_dimensions[get_column_letter(i)].width = min(max_len + 3, 55)

    # ── Confort de lecture : figer l'en-tête, activer les filtres ──────
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if rows:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(n_cols)}{header_row + len(rows)}"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )